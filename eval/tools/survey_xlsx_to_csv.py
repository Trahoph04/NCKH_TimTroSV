"""Chuẩn hoá file Google Form export (.xlsx) → CSV có tên cột ổn định.

Vì sao tách khỏi `import_survey.py`: container api KHÔNG có openpyxl (xem
apps/api/requirements.txt) và thêm dep chỉ để đọc 1 file là không đáng. Script
này chạy ở MÁY THẬT (openpyxl có sẵn), xuất CSV verbatim — mọi phép map sang
giá trị DB nằm ở `app.matching.import_survey` để còn test được bằng pytest.

CSV trung gian cũng là thứ cho phép chạy lại import mà không cần file gốc, và
là vết audit: nếu Google Form đổi câu hỏi, diff CSV sẽ chỉ đúng chỗ lệch.

    python3 eval/tools/survey_xlsx_to_csv.py "<đường dẫn .xlsx>" eval/data/survey_v2.csv

CẢNH BÁO RIÊNG TƯ: CSV chứa họ tên + MSSV + email SV thật. `eval/data/` đã nằm
trong .gitignore — đừng commit, đừng gửi ra ngoài.
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover - lỗi môi trường, không phải lỗi logic
    sys.exit("Thiếu openpyxl. Chạy: python3 -m pip install openpyxl")

# Thứ tự cột của form v2 (docs/planning/tao_google_form_v2.gs). Khoá ổn định ←
# vị trí cột, KHÔNG dùng tiêu đề tiếng Việt làm khoá: tiêu đề có dấu, dài, và
# Google đổi cách escape giữa các lần export.
COLUMNS: list[tuple[int, str]] = [
    (1, "timestamp"),
    (2, "email"),
    (3, "is_ctu"),
    (4, "year"),
    (5, "don_vi"),
    (6, "nganh"),
    (7, "current_housing"),
    (8, "gender"),
    (9, "search_methods"),
    (10, "biggest_difficulty"),
    (11, "bad_experiences"),
    (12, "budget"),
    (13, "plan"),
    (14, "top3_factors"),
    (15, "max_commute"),
    (16, "required_amenities"),
    (17, "roommate_gender_pref"),
    (18, "sleep_time"),
    (19, "noise_reaction"),
    (20, "dishes"),
    (21, "smoking"),
    (22, "ac_usage"),
    (23, "money_style"),
    (24, "behaviors"),
    (25, "persona"),
    (26, "deal_breakers"),
    (27, "need_from"),
    (28, "roommate_status"),
    (29, "useful_feature"),
    (30, "want_trial"),
    (31, "full_name"),
    (32, "mssv"),
]

EXPECTED_COLS = 32


def convert(xlsx_path: Path, csv_path: Path) -> int:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.worksheets[0]

    if ws.max_column < EXPECTED_COLS:
        sys.exit(
            f"File có {ws.max_column} cột, form v2 cần {EXPECTED_COLS}. "
            "Form đã đổi ⇒ cập nhật COLUMNS trước khi import."
        )

    # row_no giữ lại để truy vết về đúng dòng trong file gốc khi kiểm dữ liệu
    fieldnames = ["row_no"] + [key for _, key in COLUMNS]
    csv_path.parent.mkdir(parents=True, exist_ok=True)

    written = 0
    with csv_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        for row_no, row in enumerate(ws.iter_rows(min_row=2), start=2):
            cells = [c.value for c in row]
            if not any(v is not None for v in cells):
                continue  # dòng trống cuối sheet
            rec = {"row_no": row_no}
            for idx, key in COLUMNS:
                val = cells[idx - 1] if idx - 1 < len(cells) else None
                rec[key] = "" if val is None else str(val).strip()
            writer.writerow(rec)
            written += 1
    return written


def main() -> None:
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    src, dst = Path(sys.argv[1]), Path(sys.argv[2])
    if not src.is_file():
        sys.exit(f"Không thấy file: {src}")
    n = convert(src, dst)
    print(f"✓ {n} phản hồi → {dst}")


if __name__ == "__main__":
    main()
